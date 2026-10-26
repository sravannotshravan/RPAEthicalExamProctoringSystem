from collections import Counter
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEVERITIES = ('HIGH', 'MEDIUM', 'LOW')
HEADER_FILL = PatternFill('solid', fgColor='1F3864')
SEVERITY_FILL = {
    'HIGH': PatternFill('solid', fgColor='F8CBAD'),
    'MEDIUM': PatternFill('solid', fgColor='FFE699'),
}


def _fmt(dt):
    return dt.strftime('%Y-%m-%d %H:%M:%S') if dt else ''


def _write_sheet(ws, headers, rows, widths=None):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = 'A2'
    for idx, header in enumerate(headers, start=1):
        longest = max([len(str(header))] + [len(str(r[idx - 1])) for r in rows]) if rows else len(str(header))
        ws.column_dimensions[get_column_letter(idx)].width = min((widths or {}).get(idx, longest + 2), 80)


def build_exam_workbook(exam) -> bytes:
    """
    Builds the per-student audit workbook for an exam: scores, infraction counts and every
    logged incident. Wording stays neutral - rows are observations that need human review.
    """
    max_score = 100 * len(exam.questions)
    summary_rows, incident_rows, env_rows = [], [], []

    for session in exam.sessions:
        student = session.student
        name = student.name if student else ''
        reg = student.register_number if student else ''
        score = sum((sub.score or 0) for sub in session.submissions)
        sev = Counter(i.severity for i in session.incidents)
        types = Counter(i.incident_type for i in session.incidents)
        pending = sum(1 for i in session.incidents if i.review_status == 'PENDING')
        summary_rows.append([
            session.id, reg, name, student.email if student else '', session.status,
            _fmt(session.started_at), _fmt(session.ended_at),
            score, max_score, len(session.submissions),
            len(session.incidents), sev['HIGH'], sev['MEDIUM'], sev['LOW'],
            len(types), ', '.join(f"{t} x{n}" for t, n in types.most_common()),
            pending,
            'REQUIRES_HUMAN_REVIEW' if session.incidents else 'NO_ANOMALIES_OBSERVED',
        ])
        for inc in sorted(session.incidents, key=lambda i: i.timestamp or 0):
            incident_rows.append([
                session.id, reg, name, inc.id, _fmt(inc.timestamp), inc.incident_type,
                inc.severity, inc.source, inc.review_status, inc.details or '',
            ])
        for chk in session.environment_checks:
            env_rows.append([
                session.id, reg, name, _fmt(chk.checked_at), chk.status, chk.internet_status,
                chk.browser_detected or '', chk.webcam_status or '', chk.microphone_status or '',
                chk.prohibited_processes or '', chk.summary or '',
            ])

    wb = Workbook()
    ws = wb.active
    ws.title = 'Student Summary'
    _write_sheet(ws, [
        'SessionId', 'RegisterNumber', 'StudentName', 'Email', 'Status', 'StartedAt', 'EndedAt',
        'TotalScore', 'MaxScore', 'Submissions',
        'TotalInfractions', 'HighSeverity', 'MediumSeverity', 'LowSeverity',
        'DistinctInfractionTypes', 'InfractionBreakdown', 'PendingReview', 'Recommendation',
    ], summary_rows)

    ws = wb.create_sheet('Infractions')
    _write_sheet(ws, [
        'SessionId', 'RegisterNumber', 'StudentName', 'IncidentId', 'Timestamp', 'Type',
        'Severity', 'Source', 'ReviewStatus', 'Details',
    ], incident_rows, widths={10: 80})
    for row in ws.iter_rows(min_row=2):
        fill = SEVERITY_FILL.get(row[6].value)
        if fill:
            row[6].fill = fill

    ws = wb.create_sheet('Environment Checks')
    _write_sheet(ws, [
        'SessionId', 'RegisterNumber', 'StudentName', 'CheckedAt', 'Status', 'Internet', 'Browser',
        'Webcam', 'Microphone', 'ProhibitedProcesses', 'Summary',
    ], env_rows)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
